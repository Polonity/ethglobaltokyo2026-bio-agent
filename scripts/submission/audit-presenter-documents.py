"""Audit the current PDFs/PPTX using the bundled document runtime; render review sheets."""
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw
from pypdf import PdfReader
import pypdfium2 as pdfium
from pptx import Presentation

root = Path(__file__).resolve().parents[2]
dir = root / 'docs/submission/presenter-kit'
out = root / 'artifacts/submission-presenter-current'
out.mkdir(parents=True, exist_ok=True)
report = {'auditedAt': datetime.now(timezone.utc).isoformat(), 'pdfs': [], 'pptx': [], 'visualReview': 'Pending assistant inspection of rendered PDF pages', 'nativePowerPointRendering': False}
for name, pages in [('qa-cheatsheet-ja',4),('qa-cheatsheet-en',4),('qa-cheatsheet-ja-en',8),('explanation-ja',7),('explanation-en',7)]:
    file = dir / (name+'.pdf')
    reader = PdfReader(file)
    assert len(reader.pages) == pages, (name,len(reader.pages))
    texts = [p.extract_text() for p in reader.pages]
    assert all(len(t.strip())>60 for t in texts)
    report['pdfs'].append({'file':file.name,'pages':pages,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'textCharacters':[len(t) for t in texts]})
    if name.endswith('ja-en'):
        continue
    doc = pdfium.PdfDocument(str(file))
    cell_w,cell_h = (500,710) if name.startswith('qa-') else (650,390)
    sheet = Image.new('RGB',(cell_w*2,cell_h*((pages+1)//2)),'#dbe3dc')
    draw = ImageDraw.Draw(sheet)
    for i in range(pages):
        page=doc[i]
        full=page.render(scale=1.4).to_pil().convert('RGB')
        full.save(out / f'{name}-page-{i+1}.png')
        preview=full.copy();preview.thumbnail((cell_w-12,cell_h-32))
        x,y=(i%2)*cell_w,(i//2)*cell_h
        sheet.paste(preview,(x+6,y+6));draw.text((x+10,y+cell_h-23),f'{name} / {i+1}',fill='#153c38')
        page.close()
    sheet.save(out / (name+'-contact.jpg'),quality=94)
    doc.close()
for lang in ('ja','en'):
    file=dir / f'explanation-{lang}.pptx'
    deck=Presentation(file)
    assert len(deck.slides)==7
    assert all(len(s.notes_slide.notes_text_frame.text)>150 for s in deck.slides)
    tables = [s for s in deck.slides[-1].shapes if s.has_table]
    assert len(tables)==1 and len(tables[0].table.rows)==7
    assert tables[0].top + tables[0].height < int(6.33*914400), 'Architecture table overlaps the boundary note'
    with zipfile.ZipFile(file) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml','.rels')): ET.fromstring(z.read(name))
    report['pptx'].append({'file':file.name,'slides':7,'speakerNotes':7,'editableText':True,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
(dir / 'document-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'pdfs':len(report['pdfs']),'pdfPages':sum(p['pages'] for p in report['pdfs']),'pptx':report['pptx']},ensure_ascii=False))
