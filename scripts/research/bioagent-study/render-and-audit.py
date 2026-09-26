import json,re,zipfile,sys
import xml.etree.ElementTree as ET
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
import pypdfium2 as pdfium
from PIL import Image

out=Path(sys.argv[1]);summary=json.loads((out/'summary.json').read_text(encoding='utf-8'))
report=json.loads((out/'slide-validation.json').read_text(encoding='utf-8'))
for lang in ['ja','en']:
    base=out/f'BioAgent-research-{lang}';reader=PdfReader(base.with_suffix('.pdf'));assert len(reader.pages)==3
    texts=[p.extract_text() for p in reader.pages]
    assert '39%' in texts[0] and '56%' in texts[0] and '24,300' in texts[1]
    if lang=='ja':assert all(len(re.findall('[ぁ-んァ-ン一-龯]',t))>20 for t in texts)
    pptx=Presentation(base.with_suffix('.pptx'));assert len(pptx.slides)==3
    assert all(len(s.notes_slide.notes_text_frame.text)>200 for s in pptx.slides)
    with zipfile.ZipFile(base.with_suffix('.pptx')) as z:
        assert z.testzip() is None
        for f in z.namelist():
            if f.endswith(('.xml','.rels')):ET.fromstring(z.read(f))
    doc=pdfium.PdfDocument(str(base.with_suffix('.pdf')));contact=Image.new('RGB',(1280,2160),'white')
    for i in range(3):
        im=doc[i].render(scale=4/3).to_pil().convert('RGB');im.save(out/f'pdf-{lang}-{i+1}.png');contact.paste(im.resize((1280,720)),(0,720*i))
    contact.save(out/f'contact-{lang}.jpg',quality=93);doc.close()
    report['languages'][lang]['artifactChecks']={'pdfPages':3,'pptxSlides':3,'notes':3,'editableTextShapes':sum(sh.has_text_frame for s in pptx.slides for sh in s.shapes),'pdfTextNumbers':True,'pptxXmlAndZip':True,'pdfRenderer':'PDFium','pptxNativeOfficeRendered':False}
(out/'slide-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'slidesPassed':report['passed'],'charts':'rendered separately with Plotly','languages':list(report['languages'])}))
