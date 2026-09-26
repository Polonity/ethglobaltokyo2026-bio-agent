"""Inspect exported PDFs, speaker notes and editable PPTX packages."""
import hashlib
import json
import math
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw
from pypdf import PdfReader
import pypdfium2 as pdfium
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'artifacts' / 'uniswap-proposal-20260926'
report = json.loads((OUT / 'validation.json').read_text(encoding='utf-8'))
for lang in ('ja', 'en'):
    pdf_path = OUT / f'Uniswap-proposal-{lang}.pdf'
    reader = PdfReader(pdf_path)
    assert len(reader.pages) == 14
    texts = [p.extract_text() for p in reader.pages]
    assert all(len(s.strip()) > 40 for s in texts)
    assert '166,700' in texts[3] and '25,582,938' in texts[3]
    assert '1.896' in texts[5] and '2.387' in texts[5]
    assert '0.016' in texts[5] and '4.019' in texts[5]
    assert 'FEEDBACK.md' in ''.join(texts[2].split()) and 'V3' in texts[7]
    assert '100' in texts[4] and '0.001' in texts[4]
    assert all(959 < float(p.mediabox.width) < 961 for p in reader.pages)
    assert all(539 < float(p.mediabox.height) < 541 for p in reader.pages)
    pres = Presentation(OUT / f'Uniswap-proposal-{lang}.pptx')
    assert len(pres.slides) == 14
    notes = [s.notes_slide.notes_text_frame.text for s in pres.slides]
    assert all(len(n) > 100 for n in notes)
    editable = sum(1 for s in pres.slides for sh in s.shapes if sh.has_text_frame)
    with zipfile.ZipFile(OUT / f'Uniswap-proposal-{lang}.pptx') as z:
        assert z.testzip() is None
        xml_count = 0
        for name in z.namelist():
            if name.endswith(('.xml', '.rels')):
                ET.fromstring(z.read(name)); xml_count += 1
        combined = ''.join(z.read(f'ppt/slides/slide{i}.xml').decode() for i in range(1, 15))
        assert '[object Object]' not in combined and 'undefined' not in combined
    doc = pdfium.PdfDocument(str(pdf_path))
    sheet = Image.new('RGB', (1280, math.ceil(len(doc) / 4) * 206), '#e2e8ee')
    draw = ImageDraw.Draw(sheet)
    for i in range(len(doc)):
        pg = doc[i]
        im = pg.render(scale=1/3).to_pil().convert('RGB')
        im.thumbnail((316, 180))
        x, y = (i % 4) * 320, (i // 4) * 206
        sheet.paste(im, (x+2, y+3)); draw.text((x+10, y+184), f'{lang.upper()} / {i+1:02}', fill='#132d46')
        pg.render(scale=1.5).to_pil().convert('RGB').save(OUT / 'previews' / f'pdf-{lang}-{i+1:02}.png')
        pg.close()
    sheet.save(OUT / f'contact-sheet-{lang}.jpg', quality=92); doc.close()
    report['languages'][lang]['artifactChecks'] = {
        'pdfPages':len(reader.pages), 'pdfTextExtractsOnAllPages':True,
        'pdfNumericClaimsChecked':True, 'pdfRenderedWith':'PDFium',
        'pptxSlides':len(pres.slides), 'pptxNotes':len(notes), 'editableTextShapes':editable,
        'pptxZipAndXMLValid':True, 'parsedXMLParts':xml_count,
        'pptxNativeOfficeRendered':False,
        'note':'PDF visually rendered; PPTX structure and notes validated, not rendered in Microsoft PowerPoint.'
    }
    for prefix in ('Meeting-brief', 'Decision-sheet'):
        fn = f'{prefix}-{lang}'
        handout = PdfReader(OUT / f'{fn}.pdf')
        assert len(handout.pages) == 1
        assert len(handout.pages[0].extract_text()) > 100
        pd = pdfium.PdfDocument(str(OUT / f'{fn}.pdf'))
        pg = pd[0]; pg.render(scale=1.3).to_pil().save(OUT / 'previews' / f'pdf-{fn}.png'); pg.close(); pd.close()
        report['handouts'][fn]['pdfPages'] = 1
        report['handouts'][fn]['pdfRenderedWith'] = 'PDFium'

snapshot = json.loads((OUT / 'source-review.json').read_text(encoding='utf-8'))
changed = [p for p, digest in snapshot['files'].items() if hashlib.sha256((ROOT / p).read_bytes()).hexdigest() != digest]
report['sourceSnapshot'] = {'changedSinceReview': changed}
assert not changed, changed
report['passed'] = all(not d['overflow'] and not d['browserErrors'] and not d['brokenImages'] and d['notesVisible'] and d['navigationWorks'] for d in report['languages'].values()) and all(not d['beyondFooter'] and not d['horizontal'] for d in report['handouts'].values()) and not report['index']['missingLinks'] and not report['index']['videoPlayback']['error']
(OUT / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
print(json.dumps({'passed':report['passed'], 'checks':{lang:report['languages'][lang]['artifactChecks'] for lang in ('ja','en')}, 'sourceSnapshot':report['sourceSnapshot']}, ensure_ascii=False, indent=2))
