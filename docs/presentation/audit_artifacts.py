"""Validate final documents using the bundled Windows Python document libraries."""
import json
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw
from pypdf import PdfReader
import pypdfium2 as pdfium
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'artifacts' / 'explanation-slides-20260926'
report = json.loads((OUT / 'validation.json').read_text(encoding='utf-8'))
for lang in ('ja', 'en'):
    pdf_path = OUT / f'BioAgent-{lang}.pdf'
    reader = PdfReader(pdf_path)
    assert len(reader.pages) == 28
    all_text = [p.extract_text() for p in reader.pages]
    assert all(len(s.strip()) > 40 for s in all_text)
    assert '166,700' in all_text[5]
    assert '25,582,938' in all_text[5]
    assert '13.164' in all_text[14]
    assert '2.387' in all_text[14]
    assert '0.222' in all_text[15] and '-3.348' in all_text[15]
    assert '155.21' in all_text[16]
    assert all(float(p.mediabox.width) > 950 for p in reader.pages)
    assert all(539 < float(p.mediabox.height) < 541 for p in reader.pages)
    pptx_path = OUT / f'BioAgent-{lang}.pptx'
    pres = Presentation(pptx_path)
    assert len(pres.slides) == 28
    notes = [s.notes_slide.notes_text_frame.text for s in pres.slides]
    assert all(len(n) > 150 for n in notes)
    editable_texts = sum(1 for s in pres.slides for sh in s.shapes if sh.has_text_frame)
    with zipfile.ZipFile(pptx_path) as z:
        assert z.testzip() is None
        xml_count = 0
        for name in z.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                ET.fromstring(z.read(name)); xml_count += 1
        shapes_text = ''.join(z.read(f'ppt/slides/slide{i}.xml').decode() for i in range(1, 29))
        assert '[object Object]' not in shapes_text
        assert 'undefined' not in shapes_text
    doc = pdfium.PdfDocument(str(pdf_path))
    sheet = Image.new('RGB', (1280, 7 * 206), '#e2e8e7')
    draw = ImageDraw.Draw(sheet)
    for i in range(len(doc)):
        pg = doc[i]
        im = pg.render(scale=1 / 3).to_pil().convert('RGB')
        im.thumbnail((316, 180))
        x, y = (i % 4) * 320, (i // 4) * 206
        sheet.paste(im, (x + 2, y + 3))
        draw.text((x + 10, y + 184), f'{lang.upper()} / {i+1:02}', fill='#142d35')
        if i in (0, 4, 6, 8, 12, 14, 15, 16, 23, 27):
            pg.render(scale=1.5).to_pil().convert('RGB').save(OUT / 'previews' / f'pdf-{lang}-{i+1:02}.png')
        pg.close()
    sheet.save(OUT / f'contact-sheet-{lang}.jpg', quality=92)
    doc.close()
    report['languages'][lang]['artifactChecks'] = {
        'pdfPages': len(reader.pages), 'pdfTextExtractsOnAllPages': True,
        'pdfNumericClaimsChecked': True, 'pdfRenderedWith': 'PDFium',
        'pptxSlides': len(pres.slides), 'pptxNotes': len(notes),
        'editableTextShapes': editable_texts, 'parsedXMLParts': xml_count,
        'pptxZipAndXMLValid': True,
        'pptxNativeOfficeRendered': False,
        'note': 'PDF visually rendered; PPTX structure and notes validated, not rendered in Microsoft PowerPoint.'
    }
report['passed'] = all(not d['overflow'] and not d['browserErrors'] and not d['brokenImages'] and d['notesVisible'] and d['navigationWorks'] for d in report['languages'].values())
(OUT / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps({lang: report['languages'][lang]['artifactChecks'] for lang in ('ja', 'en')}, ensure_ascii=False, indent=2))
