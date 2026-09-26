import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
import pypdfium2 as pdfium
from PIL import Image

out=Path(sys.argv[1])
report=json.loads((out/'validation.json').read_text(encoding='utf-8'))
contact=Image.new('RGB',(1280,1440),'white')
for row,lang in enumerate(('ja','en')):
    base=out/f'BioAgent-benefits-{lang}'
    reader=PdfReader(base.with_suffix('.pdf'))
    assert len(reader.pages)==1
    text=reader.pages[0].extract_text()
    for required in ('15%','91%','9,000','BioAgent'):
        assert required in text,(lang,required)
    if lang=='ja':assert len(re.findall('[ぁ-んァ-ン一-龯]',text))>50
    p=Presentation(base.with_suffix('.pptx'))
    assert len(p.slides)==1 and len(p.slides[0].notes_slide.notes_text_frame.text)>200
    with zipfile.ZipFile(base.with_suffix('.pptx')) as z:
        assert z.testzip() is None
        for f in z.namelist():
            if f.endswith(('.xml','.rels')):ET.fromstring(z.read(f))
    doc=pdfium.PdfDocument(str(base.with_suffix('.pdf')))
    image=doc[0].render(scale=4/3).to_pil().convert('RGB')
    image.save(out/f'pdf-{lang}.png')
    contact.paste(image.resize((1280,720)),(0,row*720))
    doc.close()
    report['languages'][lang]['artifactChecks']={'pdfPages':1,'pptxSlides':1,'speakerNotes':True,'editableTextShapes':sum(sh.has_text_frame for sh in p.slides[0].shapes),'pdfTextAndNumbers':True,'xmlAndZip':True,'pdfRenderedWith':'PDFium'}
contact.save(out/'contact-ja-en.jpg',quality=93)
(out/'validation.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
