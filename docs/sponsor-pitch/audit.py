import json
import sys
import zipfile
import math
import re
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
from PIL import Image,ImageDraw
import pypdfium2 as pdfium
import xml.etree.ElementTree as ET

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'artifacts'/'sponsor-pitches-1min-20260926'
report=json.loads((OUT/'validation.json').read_text(encoding='utf-8'))
for key,prefix in [('uniswap','Uniswap-1min'),('1inch','1inch-Aqua-1min')]:
 if len(sys.argv)>1 and key!=sys.argv[1]:continue
 for lang in ['ja','en']:
  name=f'{key}-{lang}';base=OUT/key/f'{prefix}-{lang}'
  reader=PdfReader(base.with_suffix('.pdf'));assert len(reader.pages)==3
  texts=[p.extract_text() for p in reader.pages];assert all(len(t)>90 for t in texts)
  if lang=='ja':assert all(len(re.findall('[ぁ-んァ-ン]',t))>10 for t in texts)
  if key=='uniswap':assert 'BioAgent' in texts[0] and 'V3' in texts[2] and 'Router' in texts[2]
  else:assert 'Aqua' in texts[0] and 'SwapVM' in texts[1]
  pres=Presentation(base.with_suffix('.pptx'));assert len(pres.slides)==3
  notes=[s.notes_slide.notes_text_frame.text for s in pres.slides];assert all(len(n)>100 for n in notes)
  with zipfile.ZipFile(base.with_suffix('.pptx')) as z:
   assert z.testzip() is None
   for f in z.namelist():
    if f.endswith(('.xml','.rels')):ET.fromstring(z.read(f))
  doc=pdfium.PdfDocument(str(base.with_suffix('.pdf')))
  contact=Image.new('RGB',(1280,720),'#ece8eb')
  for i in range(3):
   pg=doc[i];im=pg.render(scale=1.5).to_pil().convert('RGB');im.save(OUT/'previews'/f'pdf-{name}-{i+1}.png');im.thumbnail((640,360));contact.paste(im,((i%2)*640,(i//2)*360));pg.close()
  contact.save(OUT/'previews'/f'contact-{name}.jpg',quality=92);doc.close()
  report['decks'][name]['artifactChecks']={'pdfPages':3,'pptxSlides':3,'pptxNotes':3,'editableTextShapes':sum(1 for s in pres.slides for sh in s.shapes if sh.has_text_frame),'pdfRenderedWith':'PDFium','pptxZipAndXMLValid':True,'pptxNativeOfficeRendered':False,'japaneseTextChecked':lang=='ja'}
report['passed']=all(d['slides']==3 and d['navigationWorks'] and d['notesVisible'] and not d['overflow'] and not d['browserErrors'] for d in report['decks'].values())
(OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'passed':report['passed'],'decks':{k:v['artifactChecks'] for k,v in report['decks'].items()}},indent=2))
